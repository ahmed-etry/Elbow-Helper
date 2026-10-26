"""Synthetic assessment lists use authorized rows and ordinary workbook delivery."""
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from elbow_helper.configuration.clans import CLANS
from elbow_helper.features.agent.access import AgentAccessLost
from elbow_helper.features.agent.engine.account_lists import account_index
from elbow_helper.features.agent.engine.account_list_spreadsheets import _sheet_name
from elbow_helper.features.agent.engine.assessment import decision_answer, deliver_decision_answer
from elbow_helper.features.agent.files.delivery import delivery_manifest, publish_spreadsheet
from elbow_helper.features.agent.files.spreadsheet_tools import prepare_spreadsheet
from elbow_helper.features.agent.knowledge.answers import scoped_answer
from elbow_helper.features.agent.models import AgentAttachment
from features.agent.files.test_agent_spreadsheets import _context
from features.agent.engine import test_request_cost as cost


def columns():
    return [{'heading': 'Synthetic label', 'fact': 'name'},
            {'heading': 'Synthetic level', 'fact': 'town_hall'},
            {'heading': 'Synthetic score', 'value': 'activity'},
            {'heading': 'Synthetic explanation', 'value': 'note'},
            {'heading': 'Synthetic qualification', 'fact': 'flags'}]


def decision(**overrides):
    return json.dumps({'intent': 'Compare synthetic accounts', 'intent_source': 'request',
        'discussion_message_ids': [], 'settled_choices': [], 'open_choices': [],
        'facts': 'Synthetic estimates.\n[[accounts:first]]', 'question': '',
        'recommendation': '[[accounts:second]]',
        'account_lists_spreadsheet': 'Synthetic workbook', 'account_lists': [
            {'id': 'first', 'title': 'Group A', 'current_clan': 'BEC', 'columns': columns(), 'accounts': [
                {'account': '#S0', 'values': {'activity': '75%', 'note': 'Synthetic note'}},
                {'account': '#S1', 'values': {'activity': '12.5'}},
                {'account': 'Synthetic A', 'values': {'activity': 'ignored duplicate'}}]},
            {'id': 'second', 'title': 'Group B', 'current_clan': 'BE1', 'columns': columns(), 'accounts': [
                {'account': '#S2', 'values': {'activity': '~30%'}},
                {'account': '#missing', 'values': {'activity': 'unknown', 'note': 'Synthetic gap'}}]}],
        **overrides})


def index():
    return account_index({'current': {'rows': [
        {'player_tag': '#S0', 'player_name': 'Synthetic A', 'townHallLevel': 15,
         'basis': 'current_in_game', 'clan_code': 'BEC'},
        {'player_tag': '#S1', 'player_name': 'Synthetic B', 'townHallLevel': 14,
         'basis': 'current_in_game', 'clan_code': 'BEM'}]},
        'history': {'rows': [{'player_tag': '#S2', 'player_name': 'Synthetic C',
                              'townhall': 13, 'clan_code': 'BE1'}]}},
        {'unverified_selection': [{'player_tag': '#S0'}]})


class SheetNameTests(unittest.TestCase):
    def test_short_titles_are_preserved_and_long_titles_end_at_words(self):
        for title, expected in (
                ('Synthetic group (tag)', 'Synthetic group (tag)'),
                ('Synthetic group (a very long synthetic tag)', 'Synthetic group'),
                ('Synthetic group (one) (another very long tag)', 'Synthetic group (one)'),
                ('Synthetic account overview for review', 'Synthetic account overview for'),
                ('Synthetic account overview now later', 'Synthetic account overview now'),
                ('Synthetic\taccount overview for review', 'Synthetic\taccount overview for'),
                ('X' * 31, 'X' * 31),
                ('X' * 40, 'X' * 31),
                (" 'Synthetic/Group' ", 'Synthetic_Group'),
                ('', 'synthetic_id'),
                ("'''", 'synthetic_id')):
            with self.subTest(title=title):
                declaration = {'id': 'synthetic_id', 'title': title}
                self.assertEqual(_sheet_name(declaration, set()), expected)
                self.assertEqual(declaration['title'], title)

    def test_suffixes_preserve_words_and_uniqueness_after_shortening(self):
        used = set()
        names = [_sheet_name({'id': f'list_{number}', 'title': title}, used)
                 for number, title in enumerate([
                     'Synthetic account overview now later',
                     'SYNTHETIC ACCOUNT OVERVIEW NOW LATER',
                     'Synthetic account overview_2',
                     *['Synthetic account overview now later'] * 10])]
        self.assertEqual(names[:3], ['Synthetic account overview now',
                                    'SYNTHETIC ACCOUNT OVERVIEW_2',
                                    'Synthetic account overview_2_2'])
        self.assertEqual(names[-1], 'Synthetic account overview_12')
        self.assertEqual(len(set(name.casefold() for name in names)), len(names))
        self.assertTrue(all(len(name) <= 31 for name in names))


class AssessmentSpreadsheetTests(unittest.IsolatedAsyncioTestCase):
    async def test_long_tab_titles_survive_workbook_delivery_with_readable_unique_names(self):
        payload = json.loads(decision())
        titles = ['Synthetic group (a very long synthetic tag)',
                  'Synthetic group (another very long tag)']
        for group, title in zip(payload['account_lists'], titles):
            group['title'] = title
        payload['account_lists_spreadsheet'] = titles[0]
        with TemporaryDirectory() as directory:
            context = _context(directory)
            await deliver_decision_answer(context, json.dumps(payload), set(), accounts=index())
            attachment, = context.state.attachments
            self.assertEqual(attachment.spreadsheet_title, titles[0])
            workbook = load_workbook(BytesIO(attachment.data))
            try:
                self.assertEqual(workbook.sheetnames, ['Synthetic group', 'Synthetic group_2'])
                for sheet in workbook:
                    self.assertEqual(tuple(cell.value for cell in sheet[1]),
                                     tuple(column['heading'] for column in columns()))
            finally:
                workbook.close()
            self.assertEqual(attachment.sheet_rows, (('Synthetic group', 2), ('Synthetic group_2', 2)))

    async def test_all_resolved_rows_in_their_tabs_and_absent_from_reply_text(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            answer = await deliver_decision_answer(context, decision(), set(), accounts=index())
            self.assertEqual(answer, 'Synthetic estimates.')
            self.assertNotIn('Synthetic A', answer)
            self.assertNotIn('[[accounts:', answer)
            self.assertEqual(len(context.state.attachments), 1)
            attachment = context.state.attachments[0]
            workbook = load_workbook(BytesIO(attachment.data))
            try:
                self.assertEqual(workbook.sheetnames, ['Group A', 'Group B'])
                self.assertEqual(list(workbook['Group A'].values), [
                    tuple(column['heading'] for column in columns()),
                    ('Synthetic A', 15, '75%', 'Synthetic note', 'owner unverified'),
                    ('Synthetic B', 14, 12.5, '', 'currently in BEM')])
                self.assertEqual(list(workbook['Group B'].values)[1:], [
                    ('Synthetic C', 13, '~30%', '', 'current clan unknown'),
                    ('#missing', '', 'unknown', 'Synthetic gap', "couldn't verify this account")])
                self.assertEqual(workbook['Group A']['C3'].data_type, 'n')
            finally:
                workbook.close()
            manifest = delivery_manifest(context.state.attachments)['files'][0]
            self.assertEqual(manifest['total_rows'], 4)
            self.assertEqual(manifest['sheets'], [
                {'name': 'Group A', 'rows': 2}, {'name': 'Group B', 'rows': 2}])
            self.assertEqual(list(Path(directory).iterdir()), [])

    async def test_complete_long_lists_and_all_twelve_tabs_replace_previous_workbook(self):
        lists = [{'id': f'list_{number}', 'title': 'Repeated/Title' * 4, 'columns': columns(), 'accounts': [
            {'account': f'#T{number}_{row}', 'values': {'activity': str(row)}} for row in range(130)]}
            for number in range(12)]
        accounts = {item['account']: {'name': item['account'], 'town_hall': 10}
                    for group in lists for item in group['accounts']}
        with TemporaryDirectory() as directory:
            context = _context(directory)
            feature = AgentAttachment('synthetic.txt', b'synthetic')
            context.state.attachments.append(feature)
            await prepare_spreadsheet(context, {'title': 'Synthetic draft', 'sheets': [
                {'name': 'Draft', 'columns': ['Value'], 'rows': [['synthetic']]}]})
            content = decision(account_lists=lists, facts='Synthetic basis.', recommendation='')
            answer = await deliver_decision_answer(context, content, set(), accounts=accounts)
            self.assertEqual(answer, 'Synthetic basis.')
            self.assertEqual(len(context.state.attachments), 2)
            self.assertIs(context.state.attachments[0], feature)
            workbook = load_workbook(BytesIO(context.state.attachments[1].data))
            try:
                self.assertEqual(len(workbook.sheetnames), 12)
                self.assertEqual(len(set(name.casefold() for name in workbook.sheetnames)), 12)
                for sheet, declaration in zip(workbook, lists):
                    self.assertLessEqual(len(sheet.title), 31)
                    self.assertEqual([row[0] for row in list(sheet.values)[1:]],
                                     [item['account'] for item in declaration['accounts']])
            finally:
                workbook.close()

    async def test_no_workbook_choice_preserves_exact_inline_output(self):
        content = json.loads(decision())
        del content['account_lists_spreadsheet']
        content = json.dumps(content)
        with TemporaryDirectory() as directory:
            context = _context(directory)
            context.state.request_text = 'Synthetic request mentioning a sheet'
            answer = await deliver_decision_answer(context, content, set(), accounts=index())
            self.assertEqual(answer, decision_answer(content, set(), accounts=index()))
            self.assertIn('Group A (2)', answer)
            self.assertEqual(context.state.attachments, [])

    async def test_preparation_failures_fall_back_without_a_file_claim(self):
        for failure in ('size', 'writer', 'slots', 'title'):
            with self.subTest(failure=failure), TemporaryDirectory() as directory:
                context = _context(directory)
                content = decision()
                if failure == 'size':
                    context.guild.filesize_limit = 1
                elif failure == 'writer':
                    context.bot.workbook_writer = SimpleNamespace(
                        write=lambda *args: (_ for _ in ()).throw(OSError('Synthetic failure')))
                elif failure == 'slots':
                    context.state.attachments.extend(AgentAttachment(f'{item}.txt', b'x')
                                                     for item in range(4))
                else:
                    content = decision(account_lists_spreadsheet='\x00')
                answer = await deliver_decision_answer(context, content, set(), accounts=index())
                self.assertEqual(answer, decision_answer(content, set(), accounts=index()))
                self.assertIn('Synthetic A', answer)
                self.assertNotIn('Synthetic workbook', answer)
                self.assertNotIn('.xlsx', answer)
                self.assertFalse(any(item.spreadsheet_title for item in context.state.attachments))
                self.assertEqual(list(Path(directory).iterdir()), [])

    async def test_access_loss_before_or_during_render_aborts_without_inline_disclosure(self):
        for during_render in (False, True):
            with self.subTest(during_render=during_render), TemporaryDirectory() as directory:
                context = _context(directory)
                writer = context.bot.workbook_writer

                def revoke(path, sheets):
                    writer.write(path, sheets)
                    context.member.roles.clear()

                if during_render:
                    context.bot.workbook_writer = SimpleNamespace(write=revoke)
                else:
                    context.member.roles.clear()
                with self.assertRaises(AgentAccessLost):
                    await deliver_decision_answer(context, decision(), set(), accounts=index())
                self.assertEqual(context.state.attachments, [])
                self.assertEqual(list(Path(directory).iterdir()), [])

    async def test_private_envelope_keeps_the_useful_answer_for_both_delivery_outcomes(self):
        for succeeds in (True, False):
            with self.subTest(succeeds=succeeds), TemporaryDirectory() as directory:
                context = _context(directory)
                if not succeeds:
                    context.guild.filesize_limit = 1
                answer = await deliver_decision_answer(context,
                    decision(needs_private_evidence=True), set(), accounts=index())
                envelope = json.loads(answer)
                self.assertEqual(set(envelope), {'answer', 'needs_private_evidence'})
                self.assertTrue(envelope['needs_private_evidence'])
                self.assertEqual(scoped_answer(context, answer), envelope['answer'])
                self.assertNotIn('.xlsx', envelope['answer'])
                self.assertEqual('Synthetic A' in envelope['answer'], not succeeds)

    async def test_prepared_workbook_uses_existing_publication_and_download_fallback(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            await deliver_decision_answer(context, decision(), set(), accounts=index())
            data = context.state.attachments[0].data
            uploaded = []

            async def upload(path, title):
                uploaded.append((path.read_bytes(), title))
                return 'https://docs.google.com/spreadsheets/d/synthetic/edit', None

            context.bot.google_publisher = SimpleNamespace(upload_workbook=AsyncMock(side_effect=upload))
            attachments = await publish_spreadsheet(context, context.state.attachments)
            self.assertEqual(uploaded, [(data, 'Synthetic workbook')])
            self.assertEqual(attachments[0].data, data)
            self.assertEqual(delivery_manifest(attachments)['files'][0]['delivery'], 'link')

    async def test_flow_prepares_a_validated_final_decision_including_format_repair(self):
        for repair in (False, True):
            with self.subTest(repair=repair), TemporaryDirectory() as directory:
                context = _context(directory)
                flow, _, _, _ = cost.RequestCostTests().flow([])
                flow.context = context
                flow.assessment_requested = True
                flow.rounder.final_attempted = True
                flow.answer_accounts = index()
                flow._run = AsyncMock(return_value='invalid' if repair else decision())
                flow._repair_decision_format = AsyncMock(return_value=SimpleNamespace(content=decision()))
                answer = await flow.run()
                self.assertEqual(answer, 'Synthetic estimates.')
                self.assertEqual(len(context.state.attachments), 1)
                self.assertEqual(flow._repair_decision_format.await_count, int(repair))

    async def test_invalid_markers_are_repaired_before_any_workbook_is_prepared(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            with self.assertRaises(ValueError):
                await deliver_decision_answer(context, decision(facts='[[accounts:missing]]'),
                                              set(), accounts=index())
            self.assertEqual(context.state.attachments, [])

    async def test_each_list_chooses_its_own_order_headings_and_judgement_fields(self):
        payload = json.loads(decision())
        first, second = payload['account_lists']
        first['columns'] = [
            {'heading': 'Synthetic confidence', 'value': 'confidence'},
            {'heading': 'Synthetic identifier', 'fact': 'player_tag'},
            {'heading': 'Synthetic clan', 'fact': 'current_clan'},
            {'heading': 'Synthetic village', 'fact': 'name'},
            {'heading': 'Synthetic caveat', 'fact': 'flags'}]
        first['accounts'] = [{'account': '#S0', 'values': {'confidence': 'provisional'}}]
        second['columns'] = [{'heading': 'Synthetic TH', 'fact': 'town_hall'},
                             {'heading': 'Synthetic label', 'fact': 'name'}]
        second['accounts'] = [{'account': '#S1'}]
        content = json.dumps(payload)
        with TemporaryDirectory() as directory:
            context = _context(directory)
            await deliver_decision_answer(context, content, set(), accounts=index())
            workbook = load_workbook(BytesIO(context.state.attachments[0].data))
            try:
                self.assertEqual(list(workbook['Group A'].values), [
                    tuple(column['heading'] for column in first['columns']),
                    ('provisional', '#S0', 'BEC', 'Synthetic A', 'owner unverified')])
                self.assertEqual(list(workbook['Group B'].values), [
                    ('Synthetic TH', 'Synthetic label'), (14, 'Synthetic B (currently in BEM)')])
            finally:
                workbook.close()
            inline = decision_answer(content, set(), accounts=index())
            self.assertIn('provisional — #S0 — BEC — Synthetic A — owner unverified', inline)
            self.assertIn('14 — Synthetic B (currently in BEM)', inline)

    async def test_omitted_flag_column_keeps_qualifications_without_inventing_a_heading(self):
        payload = json.loads(decision())
        group = payload['account_lists'][0]
        group['columns'] = [{'heading': 'Synthetic label', 'fact': 'name'},
                            {'heading': 'Synthetic activity', 'value': 'activity'}]
        group['accounts'] = [{'account': '#S0', 'values': {'activity': '75%'}}]
        payload['account_lists'] = [group]
        payload['recommendation'] = ''
        content = json.dumps(payload)
        with TemporaryDirectory() as directory:
            context = _context(directory)
            await deliver_decision_answer(context, content, set(), accounts=index())
            workbook = load_workbook(BytesIO(context.state.attachments[0].data))
            try:
                self.assertEqual(list(workbook['Group A'].values), [
                    ('Synthetic label', 'Synthetic activity'), ('Synthetic A', '75% (owner unverified)')])
            finally:
                workbook.close()
            self.assertIn('Synthetic A — 75% (owner unverified)',
                          decision_answer(content, set(), accounts=index()))

    async def test_numeric_conversion_only_inspects_model_judgements(self):
        payload = json.loads(decision())
        group = payload['account_lists'][0]
        group['columns'] = [{'heading': 'Synthetic label', 'fact': 'name'},
                            {'heading': 'Synthetic TH', 'fact': 'town_hall'},
                            {'heading': 'Synthetic estimate', 'value': 'estimate'}]
        group['accounts'] = [{'account': '#S0', 'values': {'estimate': '75'}}]
        payload['account_lists'] = [group]
        payload['recommendation'] = ''
        accounts = {'#S0': {'name': '01234', 'town_hall': 12, 'current_clan': group['current_clan']}}
        from elbow_helper.features.agent.engine import account_list_spreadsheets as sheets
        with TemporaryDirectory() as directory:
            context = _context(directory)
            with patch.object(sheets, '_estimate', wraps=sheets._estimate) as convert:
                await deliver_decision_answer(context, json.dumps(payload), set(), accounts=accounts)
            convert.assert_called_once_with('75')
            workbook = load_workbook(BytesIO(context.state.attachments[0].data))
            try:
                self.assertEqual(list(workbook['Group A'].values)[1], ('01234', 12, 75))
            finally:
                workbook.close()

    async def test_forged_fact_values_and_undeclared_judgements_are_rejected_before_storage(self):
        for corruption in ('extra_fact', 'reserved_key', 'mixed_column', 'unknown_fact', 'extra_value'):
            with self.subTest(corruption=corruption), TemporaryDirectory() as directory:
                payload = json.loads(decision())
                group = payload['account_lists'][0]
                if corruption == 'extra_fact':
                    group['accounts'][0]['name'] = 'Forged synthetic name'
                elif corruption == 'reserved_key':
                    group['columns'][0] = {'heading': 'Synthetic label', 'value': 'name'}
                elif corruption == 'mixed_column':
                    group['columns'][0]['value'] = 'forgery'
                elif corruption == 'unknown_fact':
                    group['columns'][0]['fact'] = 'unread_fact'
                else:
                    group['accounts'][0]['values']['name'] = 'Forged synthetic name'
                context = _context(directory)
                with self.assertRaises(ValueError):
                    await deliver_decision_answer(context, json.dumps(payload), set(), accounts=index())
                self.assertEqual(context.state.attachments, [])

    async def test_file_only_answer_remains_empty_including_a_private_envelope(self):
        for private in (False, True):
            with self.subTest(private=private), TemporaryDirectory() as directory:
                context = _context(directory)
                content = decision(facts='[[accounts:first]]', recommendation='[[accounts:second]]',
                                   needs_private_evidence=private)
                answer = await deliver_decision_answer(context, content, set(), accounts=index())
                if private:
                    self.assertEqual(json.loads(answer), {'answer': '', 'needs_private_evidence': True})
                else:
                    self.assertEqual(answer, '')
                self.assertEqual(scoped_answer(context, answer), '')
                self.assertEqual(len(context.state.attachments), 1)

    async def test_live_profile_clans_and_punctuated_notes_match_in_workbook_and_fallback(self):
        accounts = account_index({'profiles': {'family': 'read_member_accounts', 'rows': [
            {'player_tag': '#LIVE', 'player_name': 'Synthetic live', 'townHallLevel': 15,
             'profile_status': 'ok', 'clan': {'tag': CLANS['BEC'].tag, 'name': 'Synthetic clan'}}]},
            'history': {'family': 'read_cwl_results', 'rows': [
                {'player_tag': '#LIVE', 'player_name': 'Synthetic old', 'townhall': 12, 'clan_code': 'BEE'}]}})
        content = decision(facts='[[accounts:live]]', recommendation='', account_lists=[
            {'id': 'live', 'title': 'Synthetic live accounts', 'current_clan': 'BEC', 'columns': [
                {'heading': 'Synthetic name', 'fact': 'name'},
                {'heading': 'Synthetic clan', 'fact': 'current_clan'},
                {'heading': 'Synthetic note', 'value': 'note'},
                {'heading': 'Synthetic flags', 'fact': 'flags'}], 'accounts': [
                    {'account': '#LIVE', 'values': {'note': 'Synthetic snapshot.'}}]}])
        for succeeds in (False, True):
            with self.subTest(succeeds=succeeds), TemporaryDirectory() as directory:
                context = _context(directory)
                if not succeeds:
                    context.guild.filesize_limit = 1
                answer = await deliver_decision_answer(context, content, set(), accounts=accounts)
                self.assertNotIn('current clan unknown', answer)
                if succeeds:
                    self.assertEqual(answer, '')
                    workbook = load_workbook(BytesIO(context.state.attachments[0].data))
                    try:
                        self.assertEqual(list(workbook.active.values)[1],
                                         ('Synthetic live', 'BEC', 'Synthetic snapshot.', ''))
                    finally:
                        workbook.close()
                else:
                    self.assertEqual(answer, '**Synthetic live accounts (1)**\n'
                                     'Synthetic live — BEC — Synthetic snapshot.')


if __name__ == '__main__':
    unittest.main()
