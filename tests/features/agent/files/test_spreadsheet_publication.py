"""Fake publishers deliver synthetic spreadsheets through links or exact fallback files."""

from dataclasses import replace
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from features.agent.files.test_agent_spreadsheets import _context
from elbow_helper.features.agent.access import AgentAccessLost
from elbow_helper.features.agent.delivery import AgentDeliveryMixin
from elbow_helper.features.agent.files.spreadsheet_tools import (
    prepare_spreadsheet, spreadsheet_tools,
)
from elbow_helper.features.agent.plan import capability_list
from elbow_helper.features.agent.wording import ACTION_PRIVATE_ANSWER


GOOGLE_LINK = "https://docs.google.com/spreadsheets/d/synthetic-sheet/edit"
DOWNLOAD_LINK = "https://docs.google.com/spreadsheets/d/synthetic-sheet/export?format=xlsx"


def _arguments():
    return {"title": "Synthetic export", "sheets": [{
        "name": "Items", "rows_from": [
            {"name": "Synthetic A", "value": 7}, {"name": "Synthetic B", "value": 3.125},
        ], "columns": [{"field": "name", "heading": "Name"},
                       {"field": "value", "heading": "Value"}],
    }]}


class SpreadsheetPublicationTests(unittest.IsolatedAsyncioTestCase):
    def assert_links(self, view):
        self.assertEqual([(button.label, button.url) for button in view.children], [
            ("Google Sheet", GOOGLE_LINK), ("Download", DOWNLOAD_LINK),
        ])

    def surface(self, context):
        sent = SimpleNamespace(id=502, edit=AsyncMock(), delete=AsyncMock())
        sent.edit.return_value = sent
        message = SimpleNamespace(
            id=501, author=context.member, guild=context.guild,
            channel=context.source_message.channel, mentions=(),
            reply=AsyncMock(return_value=sent), archive_reply=False,
        )
        context.bot.user = SimpleNamespace(id=999)
        surface = AgentDeliveryMixin()
        surface.bot = context.bot
        return surface, message, sent, replace(context, source_message=message)

    def test_catalogue_names_the_actual_delivery_and_warns_against_inventing_links(self):
        tool = spreadsheet_tools()[0]
        catalogue = capability_list({tool.definition.name: tool})
        self.assertIn("Google Sheet and Download buttons or an XLSX attachment", catalogue)
        self.assertIn("do not paste or invent spreadsheet links", tool.definition.description)
        self.assertIn("Put tables that belong together into one spreadsheet as separate sheets",
                      tool.definition.description)
        self.assertIn("make separate spreadsheets only when the asker wants separate files",
                      tool.definition.description)

    async def test_multiple_published_workbooks_use_exact_titles_beside_downloads(self):
        titles = ("Synthetic grouped export", "S" * 80)
        for published_count in (1, 2):
            with self.subTest(published_count=published_count), TemporaryDirectory() as directory:
                context = _context(directory)
                links = [
                    f"https://docs.google.com/spreadsheets/d/synthetic-{index}/edit"
                    for index in range(published_count)
                ]
                context.bot.google_publisher = SimpleNamespace(upload_workbook=AsyncMock(
                    side_effect=[*((link, None) for link in links), (None, "Synthetic warning")],
                ))
                for title in (*titles[:published_count], "Synthetic unavailable export"):
                    await prepare_spreadsheet(context, {**_arguments(), "title": title})
                self.assertEqual(
                    [item.spreadsheet_title for item in context.state.attachments[:-1]],
                    list(titles[:published_count]),
                )
                surface, message, _, context = self.surface(context)
                with patch("elbow_helper.features.agent.delivery.can_show",
                           AsyncMock(return_value=True)):
                    await surface.send_response(message, "Synthetic exports", None,
                                                context.state.attachments, context=context)
                buttons = message.reply.await_args.kwargs["view"].children
                self.assertEqual([(button.label, button.url) for button in buttons], [
                    pair for index, link in enumerate(links) for pair in (
                        (titles[index] if published_count > 1 else "Google Sheet", link),
                        ("Download", (
                            f"https://docs.google.com/spreadsheets/d/synthetic-{index}"
                            "/export?format=xlsx"
                        )),
                    )
                ])
                self.assertEqual([button.row for button in buttons],
                                 [row for row in range(1, published_count + 1) for _ in range(2)])
                files = message.reply.await_args.kwargs["files"]
                self.assertEqual(len(files), 1)
                self.assertIn("Synthetic warning", message.reply.await_args.args[0])
                for file in files:
                    file.close()

    async def test_publish_uploads_exact_bytes_once_and_delivers_only_link_buttons(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            uploaded = []

            async def upload(path, title):
                self.assertEqual(title, "Synthetic export")
                uploaded.append((path, path.read_bytes()))
                workbook = load_workbook(path)
                try:
                    self.assertEqual(list(workbook["Items"].values), [
                        ("Name", "Value"), ("Synthetic A", 7), ("Synthetic B", 3.125),
                    ])
                finally:
                    workbook.close()
                return GOOGLE_LINK, None

            context.bot.google_publisher = SimpleNamespace(
                upload_workbook=AsyncMock(side_effect=upload),
            )
            result = await prepare_spreadsheet(context, _arguments())
            repeated = await prepare_spreadsheet(context, _arguments())
            self.assertEqual(result, repeated)
            self.assertTrue(result["google_sheet_published"])
            self.assertFalse(result["attachment_prepared"])
            self.assertNotIn(GOOGLE_LINK, json.dumps(result))
            self.assertNotIn(DOWNLOAD_LINK, json.dumps(result))
            self.assertEqual(len(context.state.attachments), 1)
            self.assertEqual(uploaded[0][1], context.state.attachments[0].data)
            self.assertFalse(uploaded[0][0].exists())
            self.assertEqual(list(Path(directory).iterdir()), [])
            context.bot.google_publisher.upload_workbook.assert_awaited_once()

            surface, message, _, context = self.surface(context)
            with patch("elbow_helper.features.agent.delivery.can_show",
                       AsyncMock(return_value=True)):
                await surface.send_response(message, "Synthetic workbook ready", None,
                                            context.state.attachments, context=context)
            call = message.reply.await_args
            self.assertEqual(call.args[0], "Synthetic workbook ready")
            self.assertNotIn("files", call.kwargs)
            self.assert_links(call.kwargs["view"])
            message.reply.reset_mock()
            with patch("elbow_helper.features.agent.delivery.can_show",
                       AsyncMock(return_value=True)):
                await surface.send_response(message, "", None,
                                            context.state.attachments, context=context)
            self.assertIsNone(message.reply.await_args.args[0])
            self.assertNotIn("files", message.reply.await_args.kwargs)
            self.assert_links(message.reply.await_args.kwargs["view"])

    async def test_publisher_failure_attaches_the_workbook_and_preserves_its_warning(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            warning = "Couldn't create the Google Sheet."
            context.bot.google_publisher = SimpleNamespace(
                upload_workbook=AsyncMock(return_value=(None, warning)),
            )
            result = await prepare_spreadsheet(context, _arguments())
            self.assertFalse(result["google_sheet_published"])
            self.assertTrue(result["attachment_prepared"])
            self.assertEqual(await prepare_spreadsheet(context, _arguments()), result)
            context.bot.google_publisher.upload_workbook.assert_awaited_once()
            self.assertEqual(list(Path(directory).iterdir()), [])
            surface, message, _, context = self.surface(context)
            with patch("elbow_helper.features.agent.delivery.can_show",
                       AsyncMock(return_value=True)):
                await surface.send_response(message, "Synthetic workbook ready", None,
                                            context.state.attachments, context=context)
            call = message.reply.await_args
            self.assertEqual(call.args[0], "Synthetic workbook ready\n\n" + warning)
            self.assertNotIn("view", call.kwargs)
            self.assertEqual(len(call.kwargs["files"]), 1)
            self.assertEqual(call.kwargs["files"][0].fp.read(), context.state.attachments[0].data)
            workbook = load_workbook(BytesIO(context.state.attachments[0].data))
            try:
                self.assertEqual(workbook["Items"]["B3"].value, 3.125)
            finally:
                workbook.close()

    async def test_restricted_reply_keeps_links_behind_the_private_button_and_post_override(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            context.bot.google_publisher = SimpleNamespace(
                upload_workbook=AsyncMock(return_value=(GOOGLE_LINK, None)),
            )
            await prepare_spreadsheet(context, _arguments())
            surface, message, sent, context = self.surface(context)
            with patch("elbow_helper.features.agent.delivery.can_show",
                       AsyncMock(return_value=False)):
                await surface.send_response(message, "Synthetic workbook ready", None,
                                            context.state.attachments, context=context)
            self.assertEqual(message.reply.await_args.args[0], ACTION_PRIVATE_ANSWER)
            view = message.reply.await_args.kwargs["view"]
            self.addCleanup(view.stop)
            self.assertTrue(all(getattr(button, "url", None) is None for button in view.children))
            interaction = SimpleNamespace(
                id=601, user=context.member,
                response=SimpleNamespace(send_message=AsyncMock(), defer=AsyncMock()),
                followup=SimpleNamespace(send=AsyncMock()),
            )
            await view.open_result(interaction)
            call = interaction.response.send_message.await_args
            self.assertTrue(call.kwargs["ephemeral"])
            self.assertNotIn("files", call.kwargs)
            self.assert_links(call.kwargs["view"])
            await view.post_here(interaction)
            self.assert_links(sent.edit.await_args.kwargs["view"])
            self.assertEqual(sent.edit.await_args.kwargs["attachments"], [])
            context.bot.google_publisher.upload_workbook.assert_awaited_once()

    async def test_size_and_access_failures_do_not_publish(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            context.bot.google_publisher = SimpleNamespace(upload_workbook=AsyncMock())
            context.guild.filesize_limit = 1
            result = await prepare_spreadsheet(context, _arguments())
            self.assertIn("attachment size limit", result["error"])
            context.bot.google_publisher.upload_workbook.assert_not_called()
            self.assertEqual(list(Path(directory).iterdir()), [])
            context.guild.filesize_limit = 8 * 1024 * 1024
            with patch(
                "elbow_helper.features.agent.files.spreadsheet_tools.require_evidence_access",
                AsyncMock(side_effect=[None, AgentAccessLost("Synthetic access loss")]),
            ):
                with self.assertRaises(AgentAccessLost):
                    await prepare_spreadsheet(context, _arguments())
            context.bot.google_publisher.upload_workbook.assert_not_called()
            self.assertEqual(context.state.attachments, [])
            self.assertEqual(list(Path(directory).iterdir()), [])

    async def test_access_loss_during_publication_queues_no_file_or_link(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)

            async def upload(path, title):
                self.assertTrue(path.exists())
                context.member.roles.clear()
                return GOOGLE_LINK, None

            context.bot.google_publisher = SimpleNamespace(
                upload_workbook=AsyncMock(side_effect=upload),
            )
            with self.assertRaises(AgentAccessLost):
                await prepare_spreadsheet(context, _arguments())
            context.bot.google_publisher.upload_workbook.assert_awaited_once()
            self.assertEqual(context.state.attachments, [])
            self.assertEqual(list(Path(directory).iterdir()), [])
