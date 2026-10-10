"""Earlier result records become complete, typed workbook rows."""

from io import BytesIO
from dataclasses import replace
from datetime import datetime, timezone
import json
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from features.agent.files.test_agent_spreadsheets import _context
from features.agent.engine.test_agent_plan_flow import _Session, _Model, _model_step
from elbow_helper.features.agent.engine.registry import build_agent_tools
from elbow_helper.features.agent.engine.service import AgentService
from elbow_helper.features.agent.engine.capability_contract import CapabilityContract
from elbow_helper.features.agent.files.spreadsheet_tools import (
    prepare_spreadsheet, spreadsheet_tools,
)
from elbow_helper.features.agent.models import RegisteredAgentTool
from elbow_helper.infrastructure.ai import AgentToolDefinition, AgentStep, AgentUsage
from elbow_helper.features.agent.plan import capability_list
from elbow_helper.features.agent.plan.checker import check_plan, valid_arguments
from elbow_helper.features.agent.plan.executor import execute_plan, resolve_arguments


def _arguments(records):
    return {"title": "Synthetic", "sheets": [{
        "name": "Scores", "rows_from": records, "columns": [
            {"field": "account.name", "heading": "Account"},
            {"field": "score", "heading": "Score"},
            {"field": "missing.value", "heading": "Missing"},
        ],
    }]}


class ResultSpreadsheetTests(unittest.IsolatedAsyncioTestCase):
    def test_catalogue_shows_result_sheet_fields(self):
        catalogue = capability_list(build_agent_tools())
        self.assertIn("rows_from:reference(list of records)", catalogue)
        self.assertIn("columns:[string]", catalogue)
        self.assertIn("columns:[{field:string,heading:string}]", catalogue)
        self.assertNotIn("columns:[string/{field:string,heading:string}]", catalogue)

    async def test_wrong_result_columns_are_refused_before_the_handler_runs(self):
        read = AsyncMock(return_value={"items": [{"value": 7}]})
        write = AsyncMock()
        spreadsheet = spreadsheet_tools()[0]
        registry = {
            "read_items": RegisteredAgentTool(AgentToolDefinition(
                "read_items", "Read synthetic items.",
                {"type": "object", "properties": {}, "additionalProperties": False},
            ), read),
            "write_table": replace(spreadsheet, handler=write, definition=replace(
                spreadsheet.definition, name="write_table",
            )),
        }
        reference = {"step": "read", "path": ["items"]}
        plan = {"goal": "Synthetic table", "effort": "low", "output": "write_table", "steps": [
            {"id": "read", "capability": "read_items", "arguments": {}, "depends_on": []},
            {"id": "write", "capability": "write_table", "depends_on": ["read"],
             "arguments": {"title": "Synthetic", "sheets": [{
                 "name": "Items", "rows_from": reference, "columns": ["Value"],
             }]}},
        ]}
        check = check_plan(plan, registry)
        self.assertTrue(check.ok)
        self.assertEqual(set(check.step_errors), {"write"})
        self.assertIn("Invalid arguments", check.step_errors["write"])
        read.assert_not_called()
        write.assert_not_called()

        async def run(step, arguments, results):
            return await registry[step["capability"]].handler(None, arguments)

        results = await execute_plan(plan, run, step_errors=check.step_errors)
        self.assertEqual(results["write"], {"error": check.step_errors["write"]})
        read.assert_awaited_once()
        write.assert_not_called()

    def test_sheet_kinds_require_their_own_fields_and_column_shapes(self):
        schema = spreadsheet_tools()[0].definition.parameters
        field_columns = [{"field": "value", "heading": "Value"}]
        sheets = [
            {"name": "Written", "rows": [[""]], "columns": ["Value"]},
            {"name": "Query", "sql": "SELECT :value", "params": {"value": 7}},
            {"name": "Report", "report_id": "synthetic", "collection": "items",
             "columns": field_columns, "sheet_name": "Source"},
            {"name": "Result", "rows_from": {"step": "read", "path": ["items"]},
             "columns": field_columns},
        ]
        required_fields = (
            ("name", "rows", "columns"), ("name", "sql"),
            ("name", "report_id", "collection", "columns"), ("name", "rows_from", "columns"),
        )
        for sheet, required_fields in zip(sheets, required_fields):
            with self.subTest(sheet=sheet):
                arguments = {"title": "Synthetic", "sheets": [sheet]}
                self.assertTrue(valid_arguments(arguments, schema, {"read"}))
                for required in required_fields:
                    missing = {key: value for key, value in sheet.items() if key != required}
                    self.assertFalse(valid_arguments(
                        {"title": "Synthetic", "sheets": [missing]}, schema, {"read"},
                    ))
                mixed = {**sheet, "sql" if "sql" not in sheet else "rows": "synthetic"}
                self.assertFalse(valid_arguments(
                    {"title": "Synthetic", "sheets": [mixed]}, schema, {"read"},
                ))
                if "columns" in sheet:
                    wrong = {**sheet, "columns": (field_columns if "rows" in sheet
                                                  else ["Value"])}
                    self.assertFalse(valid_arguments(
                        {"title": "Synthetic", "sheets": [wrong]}, schema, {"read"},
                    ))

    async def test_fake_read_and_file_capabilities_write_compacted_rows_paths_exactly(self):
        records = [{"name": f"Synthetic {index}", "value": index if index % 2 else index + 0.125}
                   for index in range(12)]
        reference = {"step": "read", "path": ["items", "rows"]}
        for wrapped in (False, True):
            with self.subTest(wrapped=wrapped), TemporaryDirectory() as directory:
                context = _context(directory)
                context.source_message.id = 500
                context.source_message.created_at = datetime(2026, 1, 1, tzinfo=timezone.utc)
                context.member.display_name = "Synthetic member"
                context.guild.name = "Synthetic guild"
                read = AsyncMock(return_value={"items": records})
                spreadsheet = spreadsheet_tools()[0]
                registry = {
                    "read_items": RegisteredAgentTool(AgentToolDefinition(
                        "read_items", "Read synthetic items.",
                        {"type": "object", "properties": {}, "additionalProperties": False},
                    ), read, contract=CapabilityContract(())),
                    "write_table": replace(spreadsheet, definition=replace(
                        spreadsheet.definition, name="write_table",
                    )),
                }
                plan = {"goal": "Synthetic table", "effort": "low", "output": "write_table",
                        "steps": [
                    {"id": "read", "capability": "read_items", "arguments": {}, "depends_on": []},
                    {"id": "file", "capability": "write_table", "depends_on": [],
                     "arguments": {"title": "Synthetic", "sheets": [{
                         "name": "Items", "rows_from": [reference] if wrapped else reference,
                         "columns": [{"field": "name", "heading": "Name"},
                                     {"field": "value", "heading": "Value"}],
                     }]}},
                ]}
                session = _Session([
                    _model_step(plan), AgentStep("Synthetic workbook ready", (), AgentUsage()),
                ], [])
                with patch("elbow_helper.features.agent.engine.service.build_agent_tools",
                           return_value=registry):
                    answer = await AgentService(_Model(session)).answer(
                        question="Synthetic table", local_context="", context=context,
                    )
                self.assertEqual(answer, "Synthetic workbook ready")
                self.assertEqual(len(session.calls), 2)
                read.assert_awaited_once()
                results = json.loads(session.calls[1][0][0].content)["results"]
                self.assertEqual(results["read"]["items"]["columns"], ["name", "value"])
                self.assertEqual(len(results["read"]["items"]["rows"]), 12)
                self.assertTrue(results["file"]["attachment_prepared"])
                workbook = load_workbook(BytesIO(context.state.attachments[0].data))
                try:
                    sheet = workbook["Items"]
                    self.assertEqual(sheet.max_row, 13)
                    for index, record in enumerate(records, 2):
                        self.assertEqual(sheet.cell(index, 1).value, record["name"])
                        self.assertEqual(sheet.cell(index, 2).value, record["value"])
                        self.assertEqual(sheet.cell(index, 2).data_type, "n")
                finally:
                    workbook.close()

    async def test_nested_missing_and_numeric_values_are_written_exactly(self):
        records = [{"account": {"name": f"Synthetic {i}"}, "score": i + 0.125}
                   for i in range(150)]
        records[1]["account"] = None
        records[2]["score"] = 7
        with TemporaryDirectory() as directory:
            context = _context(directory)
            result = await prepare_spreadsheet(context, _arguments(records))
            self.assertEqual(result["rows"], 150)
            workbook = load_workbook(BytesIO(context.state.attachments[0].data))
            try:
                sheet = workbook["Scores"]
                self.assertEqual(sheet["A2"].value, "Synthetic 0")
                self.assertEqual(sheet["A3"].value, "")
                self.assertEqual(sheet["B2"].value, 0.125)
                self.assertEqual(sheet["B4"].value, 7)
                self.assertEqual(sheet["B151"].value, 149.125)
                self.assertEqual(sheet["B2"].data_type, "n")
                self.assertEqual(sheet["C2"].value, "")
            finally:
                workbook.close()

    async def test_bad_record_shape_returns_specific_error(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            for records in (None, {}, "bad", [1], [{}, []]):
                with self.subTest(records=records):
                    result = await prepare_spreadsheet(context, _arguments(records))
                    self.assertEqual(result, {"error": "rows_from must be a list of objects."})
            self.assertEqual(context.state.attachments, [])

    async def test_bad_referenced_shape_keeps_specific_error_for_both_reference_forms(self):
        schema = build_agent_tools()["prepare_spreadsheet"].definition.parameters
        reference = {"step": "scores", "path": ["records"]}
        with TemporaryDirectory() as directory:
            context = _context(directory)
            for records in (None, {}, {"score": 7}, "bad", [1], [[{}]]):
                for source in (reference, [reference]):
                    with self.subTest(records=records, source=source):
                        arguments = resolve_arguments(
                            _arguments(source), {"scores": {"records": records}}, schema=schema,
                        )
                        self.assertTrue(valid_arguments(arguments, schema, {"scores"}))
                        result = await prepare_spreadsheet(context, arguments)
                        self.assertEqual(result, {
                            "error": "rows_from must be a list of objects.",
                        })
            self.assertEqual(context.state.attachments, [])

    async def test_result_sheets_enforce_query_caps_without_truncation(self):
        with TemporaryDirectory() as directory:
            for constant, cap in (("MAX_DATA_BACKED_CELLS", 8),
                                  ("MAX_DATA_BACKED_CHARACTERS", 30)):
                context = _context(directory)
                with self.subTest(constant=constant), patch(
                    "elbow_helper.features.agent.files.spreadsheet_tools." + constant, cap,
                ):
                    result = await prepare_spreadsheet(context, _arguments([
                        {"account": {"name": "Synthetic"}, "score": 7} for _ in range(3)
                    ]))
                    self.assertEqual(result, {
                        "error": "The export is too large; narrow the query.",
                    })
                    self.assertEqual(context.state.attachments, [])
