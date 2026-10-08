"""SQL exports keep every selected row or fail their cap."""
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch
from contextlib import closing
from io import BytesIO
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from features.agent.files.test_agent_spreadsheets import _context
from elbow_helper.features.agent.datasets.sources import DataSource
from elbow_helper.features.agent.files.spreadsheet_tools import prepare_spreadsheet


class QuerySpreadsheetTests(unittest.IsolatedAsyncioTestCase):
    async def test_query_and_written_sheets_render_complete_literal_evidence(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            with closing(sqlite3.connect(root / "accounts.sqlite3")) as connection:
                connection.execute("CREATE TABLE accounts (id INTEGER, note TEXT)")
                connection.executemany(
                    "INSERT INTO accounts VALUES (?,?)", [(i, "=1+1") for i in range(150)],
                )
                connection.commit()
            context = _context(directory)
            context.bot.paths = SimpleNamespace(project_root=root, data_root=root)
            sources = (DataSource(
                "accounts", Path("accounts.sqlite3"),
                tables=(("accounts", "none", "Synthetic accounts"),),
            ),)
            with patch("elbow_helper.features.agent.datasets.catalogue.SOURCES", sources):
                result = await prepare_spreadsheet(context, {"title": "Synthetic", "sheets": [
                    {
                        "name": "Evidence",
                        "sql": "SELECT id,note FROM accounts.accounts ORDER BY id",
                    },
                    {"name": "Notes", "columns": ["Finding"], "rows": [["Synthetic finding"]]},
                ]})
            self.assertEqual(result["rows"], 151)
            workbook = load_workbook(BytesIO(context.state.attachments[0].data))
            try:
                self.assertEqual(workbook["Evidence"].max_row, 151)
                self.assertEqual(workbook["Evidence"]["A151"].value, "149")
                self.assertEqual(workbook["Evidence"]["B151"].data_type, "s")
                self.assertEqual(workbook["Notes"]["A2"].value, "Synthetic finding")
            finally:
                workbook.close()

    async def test_query_sheets_do_not_use_the_written_row_limit(self):
        result = {"rows": [{"Account": f"Synthetic {i}"} for i in range(150)], "truncated": False}
        with (
            patch(
                "elbow_helper.features.agent.files.spreadsheet_tools.require_evidence_access",
                AsyncMock(),
            ),
            patch(
                "elbow_helper.features.agent.files.spreadsheet_tools.query_context",
                AsyncMock(return_value=result),
            ) as query,
            patch(
                "elbow_helper.features.agent.files.spreadsheet_tools._store_spreadsheet",
                AsyncMock(return_value={}),
            ) as store,
        ):
            await prepare_spreadsheet(SimpleNamespace(), {"title": "Synthetic", "sheets": [{
                "name": "Accounts", "sql": "SELECT name AS Account FROM first.accounts",
            }]})
            self.assertEqual(len(store.await_args.args[1].sheets[0].rows), 150)
            self.assertIsNone(query.await_args.kwargs["max_rows"])
            mixed = await prepare_spreadsheet(SimpleNamespace(), {"title": "Synthetic", "sheets": [
                {"name": "Mixed", "sql": "SELECT 1", "rows": []},
            ]})
            self.assertIn("error", mixed)
            with patch(
                "elbow_helper.features.agent.files.spreadsheet_tools.MAX_DATA_BACKED_CELLS", 100,
            ):
                limited = await prepare_spreadsheet(
                    SimpleNamespace(), {"title": "Synthetic", "sheets": [
                        {"name": "Accounts", "sql": "SELECT 1"},
                    ]},
                )
                self.assertEqual(limited, {"error": "The export is too large; narrow the query."})
