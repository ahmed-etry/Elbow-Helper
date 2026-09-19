from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from elbow_helper.configuration.roles import CORE
from elbow_helper.features.agent.access import AgentAccessLost
from elbow_helper.features.agent.models import AgentAttachment, AgentRequestContext
from elbow_helper.features.agent.files.spreadsheets import parse_agent_spreadsheet
from elbow_helper.features.agent.tools import build_agent_tools
from elbow_helper.features.agent.tools.spreadsheets import prepare_spreadsheet
from elbow_helper.infrastructure.exports import LocalExportStore, WorkbookWriter


def _arguments(title="Leadership Review"):
    return {
        "title": title,
        "sheets": [{
            "name": "Accounts",
            "columns": ["Member ID", "Account", "Note"],
            "rows": [["123456789012345678", "#P0", '=HYPERLINK("x")']],
        }],
    }


def _context(export_directory: str):
    member = SimpleNamespace(id=42, roles=[SimpleNamespace(id=next(iter(CORE)))])
    guild = SimpleNamespace(
        id=1, me=member, get_member=lambda _: member,
        filesize_limit=8 * 1024 * 1024,
    )
    channel = SimpleNamespace(
        id=100, guild=guild,
        permissions_for=lambda _: SimpleNamespace(
            view_channel=True, read_message_history=True,
        ),
    )
    guild.get_channel_or_thread = lambda value: channel if value == 100 else None
    return AgentRequestContext(
        bot=SimpleNamespace(
            fetch_channel=AsyncMock(return_value=None),
            local_exports=LocalExportStore(Path(export_directory)),
            workbook_writer=WorkbookWriter(),
        ),
        guild=guild, member=member,
        source_message=SimpleNamespace(channel=channel),
        account_links=None, clan_health=None, message_search=None,
    )


class AgentSpreadsheetContractTests(unittest.TestCase):
    def test_content_drives_stable_filename_and_fingerprint(self):
        first = parse_agent_spreadsheet(_arguments("Leadership Å Review"))
        second = parse_agent_spreadsheet(_arguments("Leadership Å Review"))
        self.assertEqual(first.filename, "leadership-a-review.xlsx")
        self.assertEqual(first.content_fingerprint, second.content_fingerprint)
        self.assertEqual(first.workbook()[0][0], "Accounts")

    def test_rejects_ambiguous_or_malformed_tables(self):
        invalid = (
            {"title": " ", "sheets": _arguments()["sheets"]},
            {"title": "Report", "sheets": []},
            {"title": "Report", "sheets": [{
                "name": "Bad/Name", "columns": ["A"], "rows": [],
            }]},
            {"title": "Report", "sheets": [{
                "name": "Data", "columns": ["A", "a"], "rows": [],
            }]},
            {"title": "Report", "sheets": [{
                "name": "Data", "columns": ["A", "B"], "rows": [["x"]],
            }]},
        )
        for arguments in invalid:
            with self.subTest(arguments=arguments), self.assertRaises(ValueError):
                parse_agent_spreadsheet(arguments)

    def test_tool_schema_is_registered_with_bounded_literal_cells(self):
        tool = build_agent_tools()["prepare_spreadsheet"].definition
        self.assertEqual(tool.parameters["properties"]["sheets"]["maxItems"], 4)
        self.assertIn("literal text cells", tool.description)


class AgentSpreadsheetToolTests(unittest.IsolatedAsyncioTestCase):
    async def test_requested_workbook_is_literal_bounded_and_reused(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            first = await prepare_spreadsheet(context, _arguments())
            self.assertEqual(first["filename"], "leadership-review.xlsx")
            self.assertEqual(first["rows"], 1)
            self.assertEqual(list(Path(directory).iterdir()), [])
            attachment = context.state.attachments[0]
            workbook = load_workbook(
                BytesIO(attachment.data), read_only=True, data_only=False,
            )
            try:
                self.assertEqual(workbook.sheetnames, ["Accounts"])
                self.assertEqual(
                    workbook["Accounts"]["A2"].value, "123456789012345678",
                )
                self.assertEqual(workbook["Accounts"]["C2"].data_type, "s")
            finally:
                workbook.close()

            repeated = await prepare_spreadsheet(context, _arguments())
            self.assertEqual(repeated, first)
            self.assertEqual(len(context.state.attachments), 1)

    async def test_same_title_with_different_content_gets_content_suffix(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            await prepare_spreadsheet(context, _arguments())
            changed = _arguments()
            changed["sheets"][0]["rows"][0][2] = "Different"
            result = await prepare_spreadsheet(context, changed)
            self.assertRegex(
                result["filename"], r"^leadership-review-[0-9a-f]{8}\.xlsx$",
            )
            self.assertEqual(len(context.state.attachments), 2)

    async def test_access_loss_after_render_queues_nothing(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            with patch(
                "elbow_helper.features.agent.tools.spreadsheets.require_evidence_access",
                new=AsyncMock(side_effect=[None, AgentAccessLost("lost")]),
            ):
                with self.assertRaises(AgentAccessLost):
                    await prepare_spreadsheet(context, _arguments())
            self.assertEqual(context.state.attachments, [])
            self.assertEqual(list(Path(directory).iterdir()), [])

    async def test_attachment_and_server_size_limits_queue_nothing(self):
        with TemporaryDirectory() as directory:
            context = _context(directory)
            context.state.attachments = [
                AgentAttachment(str(index), b"x") for index in range(4)
            ]
            result = await prepare_spreadsheet(context, _arguments())
            self.assertIn("Four report files", result["error"])

            limited = _context(directory)
            limited.guild.filesize_limit = 1
            result = await prepare_spreadsheet(limited, _arguments())
            self.assertIn("attachment size limit", result["error"])
            self.assertEqual(limited.state.attachments, [])
            self.assertEqual(list(Path(directory).iterdir()), [])


if __name__ == "__main__":
    unittest.main()
