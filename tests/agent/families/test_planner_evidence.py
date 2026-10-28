"""Planner-selected inputs survive source joins, export and the model boundary."""

from dataclasses import dataclass, replace
from io import BytesIO
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

import discord
from openpyxl import load_workbook

from elbow_helper.configuration.clans import CLANS
from elbow_helper.agent.families import identity_reads, war_reads
from elbow_helper.agent.families.access_reads import TOOLS as ACCESS_TOOLS
from elbow_helper.agent.families.model_view import family_model_view
from agent.engine.test_agent_plan_flow import _context


class PlannerEvidenceTests(unittest.IsolatedAsyncioTestCase):
    async def test_multi_clan_scoring_retains_success_when_another_clan_fails(self):
        @dataclass
        class Score:
            player_tag: str
            rank: int

        @dataclass
        class Snapshot:
            rows: tuple

        def score(clan_code, war_ids):
            if clan_code == 'BEM':
                raise ValueError('Synthetic missing history')
            return Snapshot((Score('#P0', 1), Score('#Q0', 2)))

        context = replace(_context(), cwl_queries=SimpleNamespace(ass_wars=score))
        with patch.object(war_reads, 'selected_wars', AsyncMock(return_value={'war_ids': ['war']})):
            result = await war_reads.cwl_results(context, {'metric': 'ass', 'clan_codes': ['BEH', 'BEM'],
                'season': '2026-09', 'account_tags': ['#Q0']})
        self.assertEqual(result['ranked_population'], 2)
        self.assertEqual(result['rows'], [{'player_tag': '#Q0', 'rank': 2, 'clan_code': 'BEH'}])
        self.assertEqual(result['status'], 'partial')
        self.assertEqual(result['missing_sources'][0]['clan_code'], 'BEM')

    async def test_account_role_projection_is_explicit_and_preserves_unknown_owners(self):
        context = _context()
        role = SimpleNamespace(id=CLANS['BES'].member_role_id, name='Home role', is_default=lambda: False)
        context.member.roles.append(role)
        for item in context.member.roles:
            item.name = 'Role'
            item.is_default = lambda: False
        original = context.guild.get_member
        context.guild.get_member = lambda key: original(key) if key in (41, 999) else None
        context = replace(context, account_links=SimpleNamespace(get_links_by_tags=lambda tags: {
            '#P0': {'discord_user_id': 41, 'player_tag': '#P0'},
            '#Q0': {'discord_user_id': 42, 'player_tag': '#Q0'}, '#P2': None}))
        tool = next(t for t in identity_reads.TOOLS if t.definition.name == 'read_member_accounts')
        data = {tag: {'status': 'ok', 'data': {'name': tag}} for tag in ('#P0', '#Q0', '#P2')}
        with patch.object(identity_reads, 'profiles', AsyncMock(return_value=data)):
            hidden = await tool.handler(context, {'account_tags': list(data)})
            shown = await tool.handler(context, {'account_tags': list(data), 'include_roles': True})
        self.assertTrue(all('roles' not in row for row in family_model_view(hidden)['rows']))
        rows = {row['player_tag']: row for row in family_model_view(shown)['rows']}
        self.assertIn(role.id, {r['role_id'] for r in rows['#P0']['roles']})
        self.assertEqual(rows['#Q0']['member_id'], 42)
        self.assertIsNone(rows['#Q0']['member_name'])
        self.assertIsNone(rows['#P2']['member_id'])

    async def test_multi_clan_ass_exports_all_rows_and_pages_without_attack_samples(self):
        @dataclass
        class Score:
            player_tag: str
            rank: int | None
            ass_score: float | None
            attacks: int

        @dataclass
        class Snapshot:
            rows: tuple
            profile_key: str
            attack_sample: tuple

        def score(clan_code, war_ids):
            return Snapshot((Score('#P0', 1, 2.5, 7), Score('#Q0', None, None, 0)),
                clan_code, ({'description': 'x' * 2000},) * 30)

        context = replace(_context(), cwl_queries=SimpleNamespace(ass_wars=score))
        tool = next(t for t in war_reads.TOOLS if t.definition.name == 'read_cwl_results')
        with patch.object(war_reads, 'selected_wars', AsyncMock(return_value={'war_ids': ['war']})):
            result = await tool.handler(context, {'clan_codes': ['BEH', 'BEM'], 'metric': 'ass',
                'season': '2026-09', 'delivery': 'xlsx', 'sheet_by': 'clan_code', 'limit': 1})
        self.assertEqual(result['total_rows'], 4)
        self.assertEqual(result['ranked_population'], 4)
        self.assertEqual(result['artifact_status'], 'complete')
        self.assertTrue(result['next_cursor'])
        visible = family_model_view(result)
        self.assertEqual([r['profile_key'] for r in visible['clan_results']], ['BEH', 'BEM'])
        self.assertTrue(all('attack_sample' not in r for r in visible['clan_results']))
        self.assertTrue(all(r['attack_sample'] for r in result['clan_results']))
        book = load_workbook(BytesIO(context.state.attachments[0].data), read_only=True, data_only=True)
        try:
            self.assertEqual(book.sheetnames, ['BEH', 'BEM'])
            for code in book.sheetnames:
                rows = list(book[code].values)
                self.assertEqual(len(rows), 3)
                records = [dict(zip(rows[0], row)) for row in rows[1:]]
                self.assertEqual([r['clan_code'] for r in records], [code, code])
                self.assertEqual([r['ass_score'] for r in records], [2.5, None])
                self.assertEqual([r['rank'] for r in records], [1, None])
        finally:
            book.close()

    async def test_role_matrix_uses_live_channel_permissions_and_keeps_role_only_caveat(self):
        context = _context()
        role = SimpleNamespace(id=12, name='Role A', permissions=discord.Permissions(ban_members=True))
        context.guild.roles = [role]
        channel = context.source_message.channel
        channel.name = 'Synthetic room'
        channel.permissions_for = lambda actor: discord.Permissions(
            view_channel=actor.id != role.id, read_message_history=True)
        result = await ACCESS_TOOLS[0].handler(context, {'role_names': ['Role A']})
        row = result['rows'][0]
        self.assertFalse(row['view_channel'])
        self.assertIn('ban_members', row['guild_permissions'])
        self.assertTrue(row['role_only'])
        self.assertIn('Role-only access does not establish', result['coverage'])
        self.assertIn(channel.id, context.state.source_channels)
        self.assertFalse(context.state.proposed_changes)

    async def test_role_matrix_refuses_an_unreadable_channel_without_retaining_a_report(self):
        context = _context()
        context.guild.roles = [SimpleNamespace(id=12, name='Role A')]
        with patch('elbow_helper.agent.families.access_reads.accessible_message_channel',
                   AsyncMock(return_value=None)):
            result = await ACCESS_TOOLS[0].handler(context, {'role_names': ['Role A'], 'channel_ids': [102]})
        self.assertEqual(result['cause'], 'scope_refused')
        self.assertFalse(context.state.reports)
