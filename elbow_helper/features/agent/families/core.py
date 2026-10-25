"""Immutable complete family snapshots and deterministic server-side delivery."""
from dataclasses import dataclass, field, replace
import csv
from io import BytesIO, StringIO
import json
import time
from uuid import uuid4
from datetime import datetime, timezone
from base64 import urlsafe_b64decode

from .cursor import family_cursor

from elbow_helper.infrastructure.ai import AgentToolDefinition
from elbow_helper.domain.player_tags import normalize_player_tag
from ..models import AgentAttachment, RegisteredAgentTool
from ..engine.capability_contract import CapabilityContract
from ..access import require_evidence_access
from ..reports.base import retain_report, ArtifactCapacityError
from ..reports.tools import authorize_report

COMMON = {
    'report_id': {'type': 'string', 'minLength': 1, 'maxLength': 32},
    'offset': {'type': 'integer', 'minimum': 0},
    'limit': {'type': 'integer', 'minimum': 1, 'maximum': 500},
    'delivery': {'type': 'string', 'enum': ['text', 'csv', 'xlsx'],
                 'description': 'Use text for evidence reads. Choose a file only when requested, or when the requested output rows cannot fit a message.'},
    'view': {'type': 'string', 'enum': ['rows', 'count', 'tags']},
    'cursor': {'type': 'string', 'minLength': 1, 'maxLength': 300},
}
IDS = {'type': 'array', 'minItems': 1, 'maxItems': 200,
       'uniqueItems': True, 'items': {'type': 'integer', 'minimum': 1}}
TAGS = {'type': 'array', 'minItems': 1, 'maxItems': 200,
        'uniqueItems': True, 'items': {'type': 'string', 'maxLength': 20}}
DATE = {'type': 'string', 'format': 'date-time'}


@dataclass(frozen=True, slots=True)
class FamilyReport:
    report_id: str
    guild_id: int
    family: str
    payload: str
    retained_bytes: int = field(init=False)

    def __post_init__(self):
        if (type(self.guild_id) is not int or self.guild_id <= 0
                or not isinstance(self.report_id, str) or not 1 <= len(self.report_id) <= 32
                or not isinstance(self.family, str) or not self.family.startswith('read_')):
            raise ValueError('Invalid family report identity')
        data = json.loads(self.payload)
        if (not isinstance(data, dict) or set(data) != {'rows', 'summary'}
                or not isinstance(data['rows'], list) or not isinstance(data['summary'], dict)
                or any(not isinstance(row, dict) for row in data['rows'])):
            raise ValueError('Invalid family report snapshot')
        object.__setattr__(self, 'retained_bytes', len(json.dumps(self.storage_payload(),
                           ensure_ascii=False).encode('utf-8')))

    def storage_payload(self):
        return {'report_id': self.report_id, 'guild_id': self.guild_id,
                'family': self.family, 'payload': self.payload}

    def manifest(self):
        return {'report_id': self.report_id, 'kind': 'family', 'family': self.family,
                'total_rows': len(json.loads(self.payload)['rows'])}

    def page(self, offset=0, limit=100, view='rows'):
        data = json.loads(self.payload)
        rows = data['rows']
        page = rows[offset:offset + limit]
        tags = list(dict.fromkeys(tag for row in rows for tag in (
            ([row['player_tag']] if row.get('player_tag') else []) + row.get('account_tags', []) +
            [account['player_tag'] for account in row.get('accounts', []) if account.get('player_tag')])))
        next_offset = offset + limit if offset + limit < len(rows) and view == 'rows' else None
        return {**self.manifest(), **data['summary'], 'rows': page if view == 'rows' else [],
                'returned_rows': len(page) if view == 'rows' else 0,
                'count': data['summary'].get('count', len(rows)), 'tags': tags if view == 'tags' else [],
                'offset': offset, 'next_offset': next_offset,
                'next_cursor': self.cursor(next_offset) if next_offset is not None else None,
                'complete_snapshot': True}

    def cursor(self, offset):
        summary = json.loads(self.payload)['summary']
        return family_cursor(self.report_id, summary.get('requester_id'),
                             summary.get('selection', {}), offset)

    def cursor_offset(self, cursor):
        try:
            value = json.loads(urlsafe_b64decode(cursor))
            offset = value[-1]
            if type(offset) is int and offset >= 0 and cursor == self.cursor(offset):
                return offset
        except (ValueError, TypeError, IndexError, KeyError):
            pass
        raise ValueError('Cursor does not belong to this snapshot and selection')


def decode_family(row, guild_id):
    if row['guild_id'] != guild_id:
        raise ValueError('Family report belongs to another guild')
    return FamilyReport(**{key: row[key] for key in ('report_id', 'guild_id', 'family', 'payload')})


def _cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    # Spreadsheet applications must not execute text supplied by members.
    return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value


def export_report(report, kind):
    payload = json.loads(report.payload)
    rows = payload['rows']
    selection = payload['summary'].get('selection', {})
    columns = list(dict.fromkeys(key for row in rows for key in row))
    values = [[_cell(row.get(key)) for key in columns] for row in rows]
    sheet_rows = []
    # Openpyxl silently truncates strings beyond Excel's cell limit. Check the
    # serialized, formula-escaped values before constructing any workbook.
    if kind == 'xlsx' and any(isinstance(value, str) and len(value) > 32767
                              for row in [columns, *values] for value in row):
        kind = 'csv'
    if kind == 'csv':
        stream = StringIO(newline='')
        writer = csv.writer(stream)
        writer.writerow(columns)
        writer.writerows(values)
        data = stream.getvalue().encode('utf-8-sig')
    else:
        from openpyxl import Workbook
        book = Workbook(write_only=True)
        grouping = selection.get('sheet_by')
        if grouping == 'clan_code':
            groups = list(dict.fromkeys([*selection.get('clan_codes', []),
                                        *(row.get(grouping) for row in rows)]))
            if not groups or any(not isinstance(group, str) or not group for group in groups):
                raise ValueError('Clan tabs require a clan_code for every row.')
            for group in groups:
                sheet = book.create_sheet(group)
                sheet.append(columns)
                count = 0
                for row, value in zip(rows, values):
                    if row[grouping] == group:
                        sheet.append(value)
                        count += 1
                sheet_rows.append((sheet.title, count))
        else:
            sheet = book.create_sheet('Results')
            sheet.append(columns)
            for row in values:
                sheet.append(row)
            sheet_rows.append((sheet.title, len(rows)))
        stream = BytesIO()
        book.save(stream)
        data = stream.getvalue()
    return AgentAttachment(report.family.removeprefix('read_') + '.' + kind, data,
                           report_id=report.report_id, total_rows=len(rows),
                           sheet_rows=tuple(sheet_rows))


def family_tool(name, description, properties, reader, *, required=(), contract=None, usage_note=None):
    async def handle(context, arguments):
        await require_evidence_access(context)
        offset, limit = arguments.get('offset', 0), arguments.get('limit', 500)
        if type(offset) is not int or offset < 0 or type(limit) is not int or not 1 <= limit <= 500:
            return {'error': 'Choose a valid report page.', 'cause': 'invalid_arguments'}
        identifier = arguments.get('report_id')
        if identifier:
            report = context.state.reports.get(identifier)
            if (not isinstance(report, FamilyReport) or report.guild_id != context.guild.id
                    or report.family != name or not await authorize_report(context, identifier)):
                return {'error': 'That report is no longer available here.', 'cause': 'report_unavailable'}
            if set(arguments) - set(COMMON):
                return {'error': 'Read a new snapshot to change its selection.', 'cause': 'invalid_arguments'}
            if json.loads(report.payload)['summary'].get('requester_id', context.member.id) != context.member.id:
                return {'error': 'Read a new snapshot for this request.', 'cause': 'report_unavailable'}
            if arguments.get('cursor'):
                try:
                    offset = report.cursor_offset(arguments['cursor'])
                except ValueError:
                    return {'error': 'Use the cursor returned for this report.', 'cause': 'invalid_arguments'}
        else:
            if offset or arguments.get('cursor'):
                return {'error': 'Continue with the report ID to keep the same snapshot.',
                        'cause': 'missing_report'}
            selected = {key: value for key, value in arguments.items() if key not in COMMON}
            if 'account_tags' in selected:
                values = selected['account_tags']
                tags = [normalize_player_tag(value) for value in values] if isinstance(values, list) else []
                if not tags or any(tag is None for tag in tags):
                    return {'error': 'Choose valid account tags.', 'cause': 'invalid_arguments'}
                selected['account_tags'] = sorted(set(tags))
            missing = [key for key in required if key not in selected]
            if missing:
                return {'error': 'Please specify ' + ', '.join(missing) + '.',
                        'cause': 'needs_input', 'missing_fields': missing}
            # Leave room for joins, immutable retention and model projection.
            # A slow optional live batch must return available observations
            # before the executor's hard deadline cancels the entire report.
            from ..engine.budgets import TOOL_TIMEOUT_SECONDS
            deadline = time.monotonic() + max(0, TOOL_TIMEOUT_SECONDS - 5)
            read_context = replace(context, deadline_monotonic=min(deadline, context.deadline_monotonic)
                                   if context.deadline_monotonic is not None else deadline)
            result = await reader(read_context, selected)
            await require_evidence_access(context)
            if 'error' in result:
                return result
            if not isinstance(result.get('rows'), list):
                raise ValueError('Family reader must return complete rows')
            summary = {key: value for key, value in result.items() if key != 'rows'}
            summary.update(schema_version=1, selection=selected, resolved_inputs=selected,
                requester_id=context.member.id, observed_at=datetime.now(timezone.utc).isoformat(),
                provenance={'source_channels': sorted(context.state.source_channels),
                            'required_access': sorted(context.state.required_access)})
            report = FamilyReport(uuid4().hex, context.guild.id, name, json.dumps(
                {'rows': result['rows'], 'summary': summary}, ensure_ascii=False, allow_nan=False))
            try:
                retain_report(context.state.reports, report)
            except ArtifactCapacityError:
                return {'error': 'That complete report is too large. Choose a smaller period or cohort.',
                        'cause': 'capacity'}
            context.state.source_channels.add(context.source_message.channel.id)
            # Direct family calls have the same retained provenance as checked calls.
            context.state.report_sources[report.report_id] = frozenset(context.state.source_channels)
            context.state.report_access_requirements[report.report_id] = frozenset(context.state.required_access)
        delivery = arguments.get('delivery', 'text')
        model_page = report.page(offset, limit, arguments.get('view', 'rows'))
        artifact = None
        if delivery != 'text':
            if delivery not in ('csv', 'xlsx'):
                return {'error': 'Choose text, CSV or XLSX.', 'cause': 'invalid_arguments'}
            try:
                # XLSX ZIP timestamps can change between identical exports.
                # Reuse the retained rows before comparing generated bytes.
                for item in context.state.attachments:
                    retained = context.state.reports.get(item.report_id)
                    if (item.filename == report.family.removeprefix('read_') + '.' + delivery
                            and isinstance(retained, FamilyReport) and retained.family == report.family
                            and json.loads(retained.payload)['rows'] == json.loads(report.payload)['rows']
                            and {key: json.loads(retained.payload)['summary'].get('selection', {}).get(key)
                                 for key in ('sheet_by', 'clan_codes')} ==
                                {key: json.loads(report.payload)['summary'].get('selection', {}).get(key)
                                 for key in ('sheet_by', 'clan_codes')}):
                        artifact = item
                        break
                if artifact is None:
                    artifact = export_report(report, delivery)
                existing = next((item for item in context.state.attachments
                    if item.filename == artifact.filename and item.data == artifact.data), None)
                if existing is None:
                    context.state.attachments.append(artifact)
                else:
                    artifact = existing
            except (ValueError, OSError):
                # Keep the snapshot and readable rows even if export fails.
                artifact = None
        await require_evidence_access(context)
        result = model_page
        if delivery != 'text':
            result['artifact'] = ({'report_id': report.report_id, 'filename': artifact.filename,
                                   'total_rows': report.manifest()['total_rows']} if artifact else None)
            result['artifact_status'] = 'complete' if artifact else 'unavailable'
            if artifact and delivery == 'xlsx' and artifact.filename.endswith('.csv'):
                result['artifact_note'] = 'XLSX cell limit exceeded; supplied the complete result as lossless CSV.'
        # Page complete records before the generic transport size boundary.
        # The retained report and cursor preserve every row, including when one
        # rich family needs more than one model-sized page.
        from ..engine.budgets import MAX_TOOL_RESULT_CHARACTERS
        from ..plan.results import model_result, model_view
        def model_size(value):
            return len(json.dumps(model_view(model_result(value)), ensure_ascii=False,
                                  separators=(',', ':')))
        if result['rows'] and model_size(result) > MAX_TOOL_RESULT_CHARACTERS:
            low, high, kept = 1, len(result['rows']), 0
            while low <= high:
                middle = (low + high) // 2
                page = report.page(offset, middle, arguments.get('view', 'rows'))
                candidate = {**result, **page}
                if model_size(candidate) <= MAX_TOOL_RESULT_CHARACTERS:
                    kept = middle
                    low = middle + 1
                else:
                    high = middle - 1
            if kept:
                result.update(report.page(offset, kept, arguments.get('view', 'rows')))
            elif artifact is not None:
                # A requested export can carry an indivisible oversized record.
                # The artifact contains its exact value; do not discard success
                # because the same cell cannot also fit in model evidence.
                result.update(rows=[], returned_rows=0, next_offset=None, next_cursor=None,
                              evidence_projection='complete_requested_file')
            else:
                return {'error': 'That report row is too large to read here. Choose a narrower projection.',
                        'cause': 'capacity', 'report_id': report.report_id}
        return result

    contract = contract or CapabilityContract()
    if 'cohort_ref' in properties:
        contract = replace(contract,entity_fields=(*contract.entity_fields,('cohort_ref','family_report')))
    contract = replace(contract, entity_fields=(*contract.entity_fields, ("report_id", "family_report")),
                       retained_fields=(*contract.retained_fields, "report_id"))
    return RegisteredAgentTool(AgentToolDefinition(name, description,
        {'type': 'object', 'properties': {**properties, **COMMON},
         'required': [], 'additionalProperties': False},
        usage_note=usage_note or 'Fresh selection or report_id + cursor continuation. Follow next_cursor to completion. Exports use every retained row.'),
        handle, contract=contract,
        returns='report_id,rows:[{typed family fields}],total_rows,next_offset,complete_snapshot')
