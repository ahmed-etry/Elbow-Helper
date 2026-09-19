"""Authorized Discord attachment acquisition and bounded import artifacts."""

from __future__ import annotations

import asyncio
import csv
from datetime import date, datetime, time, timedelta, timezone
import hashlib
import io
import math
from pathlib import PurePosixPath
import re
from time import monotonic
from typing import Any
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile

from openpyxl import load_workbook


from .contracts import (
    MAX_CSV_BYTES,
    MAX_CSV_ROWS,
    MAX_CSV_CELL_CHARACTERS,
    MAX_XLSX_BYTES,
    MAX_XLSX_UNCOMPRESSED_BYTES,
    MAX_XLSX_ARCHIVE_ENTRIES,
    MAX_XLSX_SHEETS,
    MAX_XLSX_ROWS,
    MAX_XLSX_COLUMNS,
    MAX_XLSX_CELLS,
    MAX_XLSX_CELL_CHARACTERS,
    MAX_XLSX_PARSE_SECONDS,
    MAX_TEXT_BYTES,
    MAX_TEXT_CHARACTERS,
    MAX_TEXT_LINES,
    MAX_TEXT_LINE_CHARACTERS,
    SUPPORTED_CSV_CONTENT_TYPES,
    DELIMITERS,
    SUPPORTED_XLSX_CONTENT_TYPES,
    SUPPORTED_TEXT_CONTENT_TYPES,
    TEXT_EXTENSIONS,
    _REQUIRED_XLSX_PARTS,
    _UNSUPPORTED_XLSX_PATH_PARTS,
    AttachmentValidationError,
    TextImportArtifact,
    CsvImportArtifact,
    XlsxSheet,
    XlsxImportArtifact,
    _validate_shape,
    _table_issues,
    _unsupported_text_control,
    _safe_filename,
)


def attachment_metadata(message: Any) -> tuple[dict[str, Any], ...]:
    result = []
    for attachment in getattr(message, "attachments", ()):
        attachment_id = getattr(attachment, "id", None)
        size = getattr(attachment, "size", None)
        if type(attachment_id) is not int or attachment_id <= 0:
            continue
        result.append({
            "attachment_id": attachment_id,
            "filename": str(getattr(attachment, "filename", "")),
            "content_type": getattr(attachment, "content_type", None),
            "size": size if type(size) is int and size >= 0 else None,
        })
    return tuple(result)


async def acquire_csv_attachment(
    message: Any, attachment_id: int, *, report_id: str,
    delimiter_name: str | None = None, imported_at: datetime | None = None,
) -> CsvImportArtifact:
    attachment = next((item for item in getattr(message, "attachments", ())
                       if getattr(item, "id", None) == attachment_id), None)
    if attachment is None:
        raise AttachmentValidationError("That attachment is not available on the authorized message")
    filename = str(getattr(attachment, "filename", ""))
    if not _safe_filename(filename) or not filename.casefold().endswith(".csv"):
        raise AttachmentValidationError("The attachment must have a safe .csv filename")
    content_type = getattr(attachment, "content_type", None)
    if content_type:
        content_type = str(content_type).split(";", 1)[0].strip().casefold()
        if content_type not in SUPPORTED_CSV_CONTENT_TYPES:
            raise AttachmentValidationError("The attachment content type is not supported for CSV")
    declared_size = getattr(attachment, "size", None)
    if type(declared_size) is not int or declared_size < 0 or declared_size > MAX_CSV_BYTES:
        raise AttachmentValidationError("The attachment size is missing or exceeds the CSV limit")
    data = await attachment.read(use_cached=False)
    if not isinstance(data, bytes):
        raise AttachmentValidationError("The downloaded attachment did not contain bytes")
    if len(data) > MAX_CSV_BYTES:
        raise AttachmentValidationError("The downloaded attachment exceeds the CSV limit")
    if declared_size != len(data):
        raise AttachmentValidationError("The downloaded attachment size does not match its metadata")
    columns, rows, delimiter, issues = await asyncio.to_thread(
        parse_csv, data, delimiter_name=delimiter_name,
    )
    now = imported_at or datetime.now(timezone.utc)
    if now.tzinfo is None:
        now = now.replace(tzinfo=timezone.utc)
    return CsvImportArtifact(
        report_id=report_id, guild_id=message.guild.id, channel_id=message.channel.id,
        message_id=message.id, attachment_id=attachment_id, filename=filename,
        content_type=content_type, byte_size=len(data), sha256=hashlib.sha256(data).hexdigest(),
        imported_at=now.astimezone(timezone.utc).isoformat(), parser_version=1,
        encoding="utf-8-sig", delimiter=delimiter, columns=columns, rows=rows,
        issues=issues,
    )


async def acquire_xlsx_attachment(
    message: Any, attachment_id: int, *, report_id: str,
    imported_at: datetime | None = None,
) -> XlsxImportArtifact:
    attachment = next((item for item in getattr(message, "attachments", ())
                       if getattr(item, "id", None) == attachment_id), None)
    if attachment is None:
        raise AttachmentValidationError("That attachment is not available on the authorized message")
    filename = str(getattr(attachment, "filename", ""))
    if not _safe_filename(filename) or not filename.casefold().endswith(".xlsx"):
        raise AttachmentValidationError("The attachment must have a safe .xlsx filename")
    content_type = getattr(attachment, "content_type", None)
    if content_type:
        content_type = str(content_type).split(";", 1)[0].strip().casefold()
        if content_type not in SUPPORTED_XLSX_CONTENT_TYPES:
            raise AttachmentValidationError("The attachment content type is not supported for XLSX")
    declared_size = getattr(attachment, "size", None)
    if type(declared_size) is not int or declared_size < 0 or declared_size > MAX_XLSX_BYTES:
        raise AttachmentValidationError("The attachment size is missing or exceeds the XLSX limit")
    data = await attachment.read(use_cached=False)
    if not isinstance(data, bytes):
        raise AttachmentValidationError("The downloaded attachment did not contain bytes")
    if len(data) > MAX_XLSX_BYTES:
        raise AttachmentValidationError("The downloaded attachment exceeds the XLSX limit")
    if declared_size != len(data):
        raise AttachmentValidationError("The downloaded attachment size does not match its metadata")
    sheets, uncompressed_size = await asyncio.to_thread(parse_xlsx, data)
    now = imported_at or datetime.now(timezone.utc)
    if now.tzinfo is None:
        now = now.replace(tzinfo=timezone.utc)
    return XlsxImportArtifact(
        report_id=report_id, guild_id=message.guild.id, channel_id=message.channel.id,
        message_id=message.id, attachment_id=attachment_id, filename=filename,
        content_type=content_type, byte_size=len(data), uncompressed_size=uncompressed_size,
        sha256=hashlib.sha256(data).hexdigest(),
        imported_at=now.astimezone(timezone.utc).isoformat(), parser_version=1,
        sheets=sheets,
    )


async def acquire_text_attachment(
    message: Any, attachment_id: int, *, report_id: str,
    imported_at: datetime | None = None,
) -> TextImportArtifact:
    attachment = next((
        item for item in getattr(message, "attachments", ())
        if getattr(item, "id", None) == attachment_id
    ), None)
    if attachment is None:
        raise AttachmentValidationError(
            "That attachment is not available on the authorized message"
        )
    filename = str(getattr(attachment, "filename", ""))
    suffix = PurePosixPath(filename.casefold()).suffix
    if not _safe_filename(filename) or suffix not in TEXT_EXTENSIONS:
        raise AttachmentValidationError(
            "The attachment must have a safe .txt, .md, or .markdown filename"
        )
    content_type = getattr(attachment, "content_type", None)
    if content_type:
        content_type = str(content_type).split(";", 1)[0].strip().casefold()
        if content_type not in SUPPORTED_TEXT_CONTENT_TYPES:
            raise AttachmentValidationError(
                "The attachment content type is not supported for text"
            )
    declared_size = getattr(attachment, "size", None)
    if (
        type(declared_size) is not int or declared_size <= 0
        or declared_size > MAX_TEXT_BYTES
    ):
        raise AttachmentValidationError(
            "The attachment size is missing, empty, or exceeds the text limit"
        )
    data = await attachment.read(use_cached=False)
    if not isinstance(data, bytes):
        raise AttachmentValidationError(
            "The downloaded attachment did not contain bytes"
        )
    if len(data) > MAX_TEXT_BYTES:
        raise AttachmentValidationError(
            "The downloaded attachment exceeds the text limit"
        )
    if declared_size != len(data):
        raise AttachmentValidationError(
            "The downloaded attachment size does not match its metadata"
        )
    text, line_count = await asyncio.to_thread(parse_text, data)
    now = imported_at or datetime.now(timezone.utc)
    if now.tzinfo is None:
        now = now.replace(tzinfo=timezone.utc)
    return TextImportArtifact(
        report_id=report_id, guild_id=message.guild.id,
        channel_id=message.channel.id, message_id=message.id,
        attachment_id=attachment_id, filename=filename,
        content_type=content_type, byte_size=len(data),
        sha256=hashlib.sha256(data).hexdigest(),
        imported_at=now.astimezone(timezone.utc).isoformat(), parser_version=1,
        encoding="utf-8-sig", document_format=TEXT_EXTENSIONS[suffix],
        text=text, line_count=line_count,
    )


def parse_text(data: bytes) -> tuple[str, int]:
    if not data or len(data) > MAX_TEXT_BYTES:
        raise AttachmentValidationError("Text is empty or exceeds the byte limit")
    if b"\x00" in data:
        raise AttachmentValidationError("Text contains NUL bytes")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise AttachmentValidationError("Text must use UTF-8 encoding") from error
    if not text or len(text) > MAX_TEXT_CHARACTERS:
        raise AttachmentValidationError(
            "Text is empty or exceeds the character limit"
        )
    if _unsupported_text_control(text):
        raise AttachmentValidationError("Text contains unsupported control characters")
    lines = text.splitlines()
    if not 1 <= len(lines) <= MAX_TEXT_LINES:
        raise AttachmentValidationError("Text exceeds the line limit")
    if any(len(line) > MAX_TEXT_LINE_CHARACTERS for line in lines):
        raise AttachmentValidationError("Text contains an oversized line")
    return text, len(lines)


def parse_csv(
    data: bytes, *, delimiter_name: str | None = None,
) -> tuple[tuple[str, ...], tuple[tuple[str, ...], ...], str, tuple[str, ...]]:
    if len(data) > MAX_CSV_BYTES:
        raise AttachmentValidationError("CSV exceeds the byte limit")
    if b"\x00" in data:
        raise AttachmentValidationError("CSV contains NUL bytes")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise AttachmentValidationError("CSV must use UTF-8 encoding") from error
    delimiter = _select_delimiter(text, delimiter_name)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    try:
        header = next(reader)
    except StopIteration as error:
        raise AttachmentValidationError("CSV is empty") from error
    except csv.Error as error:
        raise AttachmentValidationError("CSV header is malformed") from error
    columns = tuple(value.strip() for value in header)
    if not columns or any(not value for value in columns):
        raise AttachmentValidationError("CSV headings must not be empty")
    issues = []
    duplicates = sorted({value for value in columns if columns.count(value) > 1})
    if duplicates:
        issues.append("duplicate_headings:" + ",".join(duplicates))
    rows = []
    try:
        for row_number, row in enumerate(reader, start=2):
            if row_number > MAX_CSV_ROWS + 1:
                raise AttachmentValidationError("CSV exceeds the row limit")
            if len(row) != len(columns):
                raise AttachmentValidationError(f"CSV row {row_number} has {len(row)} cells; expected {len(columns)}")
            normalized = tuple(value.strip() for value in row)
            if any(len(value) > MAX_CSV_CELL_CHARACTERS for value in normalized):
                raise AttachmentValidationError(f"CSV row {row_number} contains an oversized cell")
            rows.append(normalized)
    except csv.Error as error:
        raise AttachmentValidationError("CSV data is malformed") from error
    result = tuple(rows)
    formula_like = sum(
        value.startswith(("=", "+", "-", "@"))
        for row in result for value in row if value
    )
    if formula_like:
        issues.append(f"formula_like_cells:{formula_like}")
    _validate_shape(columns, result)
    return columns, result, delimiter, tuple(issues)


def parse_xlsx(data: bytes) -> tuple[tuple[XlsxSheet, ...], int]:
    deadline = monotonic() + MAX_XLSX_PARSE_SECONDS
    if len(data) > MAX_XLSX_BYTES:
        raise AttachmentValidationError("XLSX exceeds the compressed byte limit")
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if not 1 <= len(members) <= MAX_XLSX_ARCHIVE_ENTRIES:
                raise AttachmentValidationError("XLSX exceeds the archive entry limit")
            lowered_names: set[str] = set()
            uncompressed_size = 0
            for member in members:
                _check_xlsx_deadline(deadline)
                name = member.filename
                lowered = name.casefold()
                path = PurePosixPath(name)
                if (not name or "\\" in name or name.startswith("/") or ":" in name
                        or any(part in {"", ".", ".."} for part in path.parts)
                        or lowered in lowered_names):
                    raise AttachmentValidationError("XLSX contains an unsafe or duplicate archive path")
                lowered_names.add(lowered)
                if member.flag_bits & 1:
                    raise AttachmentValidationError("Encrypted XLSX entries are not supported")
                if member.compress_type not in {ZIP_STORED, ZIP_DEFLATED}:
                    raise AttachmentValidationError("XLSX uses an unsupported compression method")
                uncompressed_size += member.file_size
                if (member.file_size > MAX_XLSX_UNCOMPRESSED_BYTES
                        or uncompressed_size > MAX_XLSX_UNCOMPRESSED_BYTES):
                    raise AttachmentValidationError("XLSX exceeds the uncompressed byte limit")
                if lowered.endswith("vbaproject.bin") or any(
                    lowered.startswith(prefix) for prefix in _UNSUPPORTED_XLSX_PATH_PARTS
                ):
                    raise AttachmentValidationError("XLSX contains unsupported executable, external, or embedded content")
            if not _REQUIRED_XLSX_PARTS <= lowered_names:
                raise AttachmentValidationError("The attachment is not a valid XLSX workbook")
            for member in members:
                if not member.filename.casefold().endswith((".xml", ".rels")):
                    continue
                payload = archive.read(member)
                lowered_payload = payload.lower()
                if b"<!doctype" in lowered_payload or b"<!entity" in lowered_payload:
                    raise AttachmentValidationError("XLSX XML declarations are not supported")
                if (b"macroenabled" in lowered_payload or b"vbaproject" in lowered_payload
                        or re.search(br"targetmode\s*=\s*['\"]external['\"]", lowered_payload)):
                    raise AttachmentValidationError("XLSX contains unsupported executable or external content")
                if (member.filename.casefold().startswith("xl/worksheets/")
                        and b"<mergecells" in lowered_payload):
                    raise AttachmentValidationError("XLSX merged cells are not supported")
            _check_xlsx_deadline(deadline)
    except (BadZipFile, OSError) as error:
        raise AttachmentValidationError("The attachment is not a valid XLSX archive") from error

    try:
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False,
        )
    except Exception as error:
        raise AttachmentValidationError("The XLSX workbook could not be parsed safely") from error
    try:
        _check_xlsx_deadline(deadline)
        if not 1 <= len(workbook.worksheets) <= MAX_XLSX_SHEETS:
            raise AttachmentValidationError("XLSX exceeds the sheet limit")
        if any(sheet.sheet_state != "visible" for sheet in workbook.worksheets):
            raise AttachmentValidationError("Hidden XLSX sheets are not supported")
        parsed = []
        total_cells = 0
        for worksheet in workbook.worksheets:
            if worksheet.max_row > MAX_XLSX_ROWS + 1 or worksheet.max_column > MAX_XLSX_COLUMNS:
                raise AttachmentValidationError("XLSX exceeds a worksheet shape limit")
            values: list[tuple[str, ...]] = []
            formula_count = 0
            for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                _check_xlsx_deadline(deadline)
                converted = []
                for cell in row:
                    if cell.data_type == "f":
                        formula_count += 1
                        converted.append("")
                    elif cell.data_type == "e":
                        raise AttachmentValidationError(
                            f"XLSX sheet {worksheet.title!r} contains an error cell"
                        )
                    else:
                        converted.append(_xlsx_cell_text(cell.value))
                values.append(tuple(converted))
                total_cells += len(converted)
                if total_cells > MAX_XLSX_CELLS:
                    raise AttachmentValidationError("XLSX exceeds the workbook cell limit")
                if row_number > MAX_XLSX_ROWS + 1:
                    raise AttachmentValidationError("XLSX exceeds the row limit")
            if formula_count:
                raise AttachmentValidationError(
                    f"XLSX sheet {worksheet.title!r} contains formulas; replace them with literal values"
                )
            while values and not any(values[-1]):
                values.pop()
            if not values:
                raise AttachmentValidationError(f"XLSX sheet {worksheet.title!r} is empty")
            columns = tuple(value.strip() for value in values[0])
            rows = tuple(tuple(value.strip() for value in row) for row in values[1:])
            if not columns or any(not value for value in columns):
                raise AttachmentValidationError(
                    f"XLSX sheet {worksheet.title!r} has an empty heading"
                )
            parsed.append(XlsxSheet(
                worksheet.title, columns, rows, _table_issues(columns, rows),
            ))
        return tuple(parsed), uncompressed_size
    except AttachmentValidationError:
        raise
    except Exception as error:
        raise AttachmentValidationError("The XLSX workbook could not be parsed safely") from error
    finally:
        workbook.close()


def _select_delimiter(text: str, delimiter_name: str | None) -> str:
    if delimiter_name is not None:
        try:
            return DELIMITERS[delimiter_name]
        except KeyError as error:
            raise AttachmentValidationError("Unsupported CSV delimiter") from error
    sample = text[:8192]
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t").delimiter
    except csv.Error:
        if all(character not in sample for character in (";", "\t")):
            return ","
        raise AttachmentValidationError("CSV delimiter is ambiguous; choose one explicitly")


def _xlsx_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            raise AttachmentValidationError("XLSX contains a non-finite numeric cell")
        return str(value)
    if isinstance(value, str):
        if len(value) > MAX_XLSX_CELL_CHARACTERS:
            raise AttachmentValidationError("XLSX contains an oversized cell")
        return value
    raise AttachmentValidationError("XLSX contains an unsupported cell value")


def _check_xlsx_deadline(deadline: float) -> None:
    if monotonic() > deadline:
        raise AttachmentValidationError("XLSX parsing exceeded the time limit")
